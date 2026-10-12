"""Persist original policy files; cache one request-reusable vector store per company."""

from __future__ import annotations

import hashlib
import io
import json
import os
import sqlite3
from bisect import bisect_right
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock, RLock
from typing import Callable
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.vectorstores import InMemoryVectorStore
from lxml.etree import XMLSyntaxError

from apps.api.services.RAG import rag

API_DIR = Path(__file__).resolve().parents[2]
SUPPORTED_TYPES = {".pdf", ".docx", ".txt"}
MAX_DOCX_EXPANDED_BYTES = 30_000_000


class UnknownCompany(LookupError):
    pass


class PolicyNotFound(LookupError):
    pass


def company_catalog() -> list[dict]:
    return json.loads((API_DIR / "companies.json").read_text(encoding="utf-8"))


def validate_company(company_id: str) -> None:
    if company_id not in {company["id"] for company in company_catalog()}:
        raise UnknownCompany("Company not found")


def clean_filename(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not name or len(name) > 255 or any(ord(c) < 32 for c in name):
        raise ValueError("Invalid policy filename")
    if Path(name).suffix.lower() not in SUPPORTED_TYPES:
        raise ValueError("Upload a PDF, DOCX or UTF-8 TXT file")
    return name


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            # Retain text in nested tables too, in the cell's document order.
            items = [
                item.text if isinstance(item, Paragraph) else _table_text(item)
                for item in cell.iter_inner_content()
            ]
            cells.append(rag.clean("\n".join(items)))
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_file(filename: str, data: bytes) -> list[Document]:
    """Common chunk format; non-PDF citations use original lines/paragraphs/tables."""
    filename = clean_filename(filename)
    if not data or len(data) > rag.MAX_POLICY_BYTES:
        raise ValueError("Upload a nonempty file no larger than 10 MB")
    suffix = Path(filename).suffix.lower()
    common = {"kind": "policy", "source": filename, "article": filename}
    if suffix == ".pdf":
        chunks = rag.extract_policy(data)
        for chunk in chunks:
            chunk.metadata.update(common, location_kind="page")
        return chunks

    chunks = []
    if suffix == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT policies must use UTF-8 encoding") from exc
        if "\x00" in text:
            raise ValueError("The uploaded TXT contains binary data")
        # Clean each line separately, preserving the original line numbering.
        text = "\n".join(rag.clean(line) for line in text.splitlines())
        starts = [0] + [i + 1 for i, char in enumerate(text) if char == "\n"]
        for number, part in enumerate(rag.SPLITTER.create_documents([text])):
            start = part.metadata["start_index"]
            end = start + len(part.page_content) - 1
            chunks.append(Document(page_content=part.page_content, metadata={
                **common, "location_kind": "line", "location_start": bisect_right(starts, start),
                "location_end": bisect_right(starts, end), "chunk": number,
            }))
    else:
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(member.file_size for member in archive.infolist()) > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("Expanded DOCX content exceeds 30 MB")
                if any(member.flag_bits & 1 for member in archive.infolist()):
                    raise ValueError("Encrypted DOCX files are unsupported")
            document = WordDocument(io.BytesIO(data))
            paragraph_no = table_no = 0
            section = ""
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    paragraph_no += 1
                    body = rag.clean(item.text)
                    if item.style and item.style.name.startswith("Heading"):
                        section = body
                    kind, position = "paragraph", paragraph_no
                else:
                    table_no += 1
                    body = rag.clean(_table_text(item))
                    # A completely empty table is not selectable policy text.
                    if not body.replace("|", "").strip():
                        body = ""
                    kind, position = "table", table_no
                for part in rag.SPLITTER.split_text(body):
                    chunks.append(Document(page_content=part, metadata={
                        **common, "location_kind": kind, "location_start": position,
                        "location_end": position, "section": section, "chunk": len(chunks),
                    }))
        except (BadZipFile, PackageNotFoundError, KeyError, XMLSyntaxError) as exc:
            raise ValueError("The uploaded DOCX is invalid or damaged") from exc
    if not chunks:
        raise ValueError("No selectable text in policy; scanned PDFs need OCR")
    return chunks


class PolicyService:
    def __init__(self, directory: Path, embedding_factory: Callable[[], Embeddings] | None = None):
        directory.mkdir(parents=True, exist_ok=True)
        self.database = directory / "policies.sqlite3"
        self._embeddings = embedding_factory or rag.embeddings
        self._stores: dict[str, InMemoryVectorStore | None] = {}
        self._lock = RLock()
        with self._connection() as db:
            db.execute("""CREATE TABLE IF NOT EXISTS policies (
                id TEXT PRIMARY KEY, company_id TEXT NOT NULL, filename TEXT NOT NULL,
                size INTEGER NOT NULL, uploaded_at TEXT NOT NULL, sha256 TEXT NOT NULL,
                data BLOB NOT NULL, UNIQUE(company_id, sha256)
            )""")

    @contextmanager
    def _connection(self):
        db = sqlite3.connect(self.database, timeout=30)
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    @staticmethod
    def _info(row) -> dict:
        return {key: row[key] for key in ("id", "filename", "size", "uploaded_at")}

    @staticmethod
    def _documents(policy_id: str, company_id: str, filename: str, data: bytes) -> list[Document]:
        docs = extract_file(filename, data)
        for number, doc in enumerate(docs):
            doc.id = f"{policy_id}-{number}"
            doc.metadata.update(policy_id=policy_id, company_id=company_id)
        return docs

    def _store(self, company_id: str) -> InMemoryVectorStore | None:
        # Called under the service lock. Publish the cache only after a full rebuild.
        if company_id not in self._stores:
            with self._connection() as db:
                rows = db.execute(
                    "SELECT * FROM policies WHERE company_id = ? ORDER BY uploaded_at, id",
                    (company_id,),
                ).fetchall()
            store = InMemoryVectorStore(self._embeddings()) if rows else None
            if store is not None:
                for row in rows:
                    store.add_documents(self._documents(row["id"], company_id, row["filename"], row["data"]))
            self._stores[company_id] = store
        return self._stores[company_id]

    def upload(self, company_id: str, filename: str, data: bytes) -> dict:
        validate_company(company_id)
        filename = clean_filename(filename)
        if not data or len(data) > rag.MAX_POLICY_BYTES:
            raise ValueError("Upload a nonempty file no larger than 10 MB")
        digest = hashlib.sha256(data).hexdigest()
        with self._lock:
            with self._connection() as db:
                existing = db.execute(
                    "SELECT * FROM policies WHERE company_id = ? AND sha256 = ?",
                    (company_id, digest),
                ).fetchone()
            if existing is not None:
                self._store(company_id)
                return self._info(existing)
            policy_id = "p_" + uuid4().hex
            docs = self._documents(policy_id, company_id, filename, data)
            store = self._store(company_id)
            pending = InMemoryVectorStore(self._embeddings())
            pending.add_documents(docs)
            info = {
                "id": policy_id, "filename": filename, "size": len(data),
                "uploaded_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            }
            # File bytes and metadata commit together, after embedding succeeds.
            with self._connection() as db:
                db.execute("INSERT INTO policies VALUES (?, ?, ?, ?, ?, ?, ?)", (
                    policy_id, company_id, filename, len(data), info["uploaded_at"], digest, data,
                ))
            if store is None:
                self._stores[company_id] = pending
            else:
                # Transfer already computed vectors; existing documents are never re-embedded.
                store.store.update(pending.store)
            return info

    def list(self, company_id: str) -> list[dict]:
        validate_company(company_id)
        with self._lock, self._connection() as db:
            rows = db.execute(
                "SELECT id, filename, size, uploaded_at FROM policies WHERE company_id = ? ORDER BY uploaded_at DESC, id",
                (company_id,),
            ).fetchall()
            return [self._info(row) for row in rows]

    def search(self, company_id: str, question: str, k: int = 3) -> list[Document]:
        validate_company(company_id)
        with self._lock:
            store = self._store(company_id)
            return store.similarity_search(question, k=k) if store is not None else []

    def delete(self, policy_id: str) -> None:
        with self._lock:
            with self._connection() as db:
                row = db.execute("SELECT company_id FROM policies WHERE id = ?", (policy_id,)).fetchone()
                if row is None:
                    raise PolicyNotFound("Policy not found")
                db.execute("DELETE FROM policies WHERE id = ?", (policy_id,))
            store = self._stores.get(row["company_id"])
            if store is not None:
                ids = [key for key, doc in store.store.items() if doc["metadata"]["policy_id"] == policy_id]
                store.delete(ids)


_service: PolicyService | None = None
_service_lock = Lock()


def get_service() -> PolicyService:
    global _service
    with _service_lock:
        if _service is None:
            directory = Path(os.environ.get("POLICIES_DATA_DIR", str(API_DIR / ".data"))).resolve()
            _service = PolicyService(directory)
        return _service
