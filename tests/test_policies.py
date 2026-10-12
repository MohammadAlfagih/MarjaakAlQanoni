"""Exercise real parsing, SQLite persistence and in-memory company isolation."""

import io
import json
import sqlite3
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from pathlib import Path
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

import pymupdf
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from langchain_core.embeddings import Embeddings

from apps.api.app.app import app
from apps.api.services.llm import chat
from apps.api.services.policies import router as routes, service


def pdf_bytes(text="Company annual leave policy: 25 paid days"):
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
        return pdf.tobytes()


def docx_bytes():
    document = WordDocument()
    document.add_heading("Leave policy", 1)
    document.add_paragraph("Employees receive annual leave.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Annual leave"
    table.cell(0, 1).text = "25 days"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class CountingEmbeddings(Embeddings):
    def __init__(self):
        self.document_calls = 0
        self.query_calls = 0
        self.fail = False

    def embed_documents(self, texts):
        self.document_calls += 1
        if self.fail:
            raise RuntimeError("Embedding failure")
        return [[1.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, text):
        self.query_calls += 1
        return [1.0, 0.0, 0.0]


class ExtractionTests(unittest.TestCase):
    def test_pdf_pages_and_filename(self):
        docs = service.extract_file("leave.pdf", pdf_bytes())
        self.assertEqual(docs[0].metadata["source"], "leave.pdf")
        self.assertEqual(docs[0].metadata["page_start"], 1)
        self.assertEqual(docs[0].metadata["location_kind"], "page")

    def test_docx_paragraphs_tables_and_no_invented_pages(self):
        docs = service.extract_file("leave.docx", docx_bytes())
        self.assertEqual({d.metadata["location_kind"] for d in docs}, {"paragraph", "table"})
        table = next(d for d in docs if d.metadata["location_kind"] == "table")
        self.assertIn("Annual leave | 25 days", table.page_content)
        self.assertEqual(table.metadata["section"], "Leave policy")
        self.assertNotIn("page_start", table.metadata)

    def test_txt_line_numbers_preserve_blank_lines(self):
        docs = service.extract_file("leave.txt", "\ufeff\n\n  الإجازة السنوية\nتحدد اللائحة أيام الإجازة\n".encode("utf-8"))
        self.assertEqual(docs[0].metadata["location_start"], 3)
        self.assertEqual(docs[0].metadata["location_end"], 4)
        self.assertIn("الإجازة السنوية", docs[0].page_content)

    def test_invalid_empty_and_oversized_files_rejected(self):
        expanded = io.BytesIO()
        with ZipFile(expanded, "w") as archive:
            archive.writestr("too-big", b"x" * 32)
        with patch.object(service, "MAX_DOCX_EXPANDED_BYTES", 16):
            with self.assertRaises(ValueError):
                service.extract_file("bad.docx", expanded.getvalue())
        for filename, data in (
            ("a.exe", b"hello"), ("a.pdf", b"not pdf"), ("a.pdf", pdf_bytes("")),
            ("a.docx", b"not zip"), ("a.txt", b"\xff"), ("a.txt", b"a\0b"),
            ("a.txt", b" \n\n"), ("a.txt", b""), ("bad\nname.txt", b"text"),
        ):
            with self.subTest(filename=filename, data=data[:10]), self.assertRaises(ValueError):
                service.extract_file(filename, data)
        with patch.object(service.rag, "MAX_POLICY_BYTES", 5):
            with self.assertRaises(ValueError):
                service.extract_file("a.txt", b"123456")


class PolicyLifecycleTests(unittest.TestCase):
    def setUp(self):
        Path("tmp").mkdir(exist_ok=True)
        self.directory = tempfile.TemporaryDirectory(dir="tmp")
        self.addCleanup(self.directory.cleanup)
        self.embeddings = CountingEmbeddings()
        self.policies = service.PolicyService(Path(self.directory.name), lambda: self.embeddings)

    def test_upload_once_search_reuses_vectors_and_second_file_keeps_first(self):
        first = self.policies.upload("c_101", "leave.txt", b"Annual leave policy")
        self.assertEqual(self.embeddings.document_calls, 1)
        duplicate = self.policies.upload("c_101", "renamed.txt", b"Annual leave policy")
        self.assertEqual(first["id"], duplicate["id"])
        self.policies.search("c_101", "annual leave")
        self.policies.search("c_101", "leave days")
        self.assertEqual(self.embeddings.document_calls, 1)
        self.assertEqual(self.embeddings.query_calls, 2)
        second = self.policies.upload("c_101", "hours.txt", b"Working hours policy")
        self.assertEqual(self.embeddings.document_calls, 2)
        results = self.policies.search("c_101", "policy")
        self.assertEqual({d.metadata["policy_id"] for d in results}, {first["id"], second["id"]})

    def test_company_isolation_with_multiple_files(self):
        first = self.policies.upload("c_101", "a.pdf", pdf_bytes())
        second = self.policies.upload("c_102", "b.docx", docx_bytes())
        for company, policy_id in (("c_101", first["id"]), ("c_102", second["id"])):
            result = self.policies.search(company, "leave")
            self.assertTrue(result)
            self.assertTrue(all(d.metadata["company_id"] == company for d in result))
            self.assertTrue(all(d.metadata["policy_id"] == policy_id for d in result))
            self.assertEqual(len(self.policies.list(company)), 1)
        calls = self.embeddings.document_calls
        self.assertEqual(self.policies.search("c_103", "leave"), [])
        self.assertEqual(self.embeddings.document_calls, calls)

    def test_uploaded_policy_reaches_both_chat_roles_before_and_after_restart(self):
        """Exercise upload -> company store -> RAG -> prompt, independent of audience."""
        http = TestClient(app)
        with patch.object(routes, "get_service", return_value=self.policies):
            upload = http.post(
                "/policies", data={"company_id": "c_101"},
                files={"file": ("leave.txt", b"Company annual leave: 25 paid days.")},
            )
        self.assertEqual(upload.status_code, 201)

        # Real policy extraction, SQLite, vector search and RAG evidence formatting.
        # Stub law storage and generation so this checks routing without downloads.
        law_store = MagicMock()
        law_store.get.return_value = {"ids": []}
        law_store.similarity_search.return_value = []
        qwen = MagicMock()
        qwen.ask.return_value = json.dumps({
            "answer": "Company annual leave is 25 paid days.",
            "evidence_refs": ["P1"], "next_steps": [],
            "conflict": False, "status": "answered",
        })
        restarted = service.PolicyService(Path(self.directory.name), lambda: self.embeddings)
        for policies in (self.policies, restarted):
            with self.subTest(restarted=policies is restarted), \
                 patch.object(chat, "get_policy_service", return_value=policies), \
                 patch.object(chat.rag, "check_index", return_value={"chunks": 0}), \
                 patch.object(chat.rag, "store", return_value=law_store), \
                 patch.object(chat, "get_client", return_value=qwen):
                prompts = []
                responses = []
                for role in ("hr", "employee"):
                    response = http.post("/chat", json={
                        "question": "كم استحق اجازة سنوية في شركتي",
                        "company_id": "c_101", "role": role, "lang": "ar", "history": [],
                    })
                    self.assertEqual(response.status_code, 200, response.text)
                    responses.append(response.json())
                    prompt = json.loads(qwen.ask.call_args.args[1])
                    self.assertNotIn("audience", prompt)
                    self.assertIn("Company annual leave: 25 paid days.", prompt["retrieved_evidence"])
                    prompts.append(prompt)
                self.assertEqual(prompts[0], prompts[1])
                self.assertEqual(responses[0]["articles"], responses[1]["articles"])
                self.assertEqual(responses[1]["articles"][0]["source"], "policy")

                calls = qwen.ask.call_count
                other_company = http.post("/chat", json={
                    "question": "How many annual leave days?", "company_id": "c_102",
                    "role": "employee", "lang": "en", "history": [],
                })
                self.assertEqual(other_company.status_code, 200)
                self.assertEqual(other_company.json()["articles"], [])
                self.assertEqual(qwen.ask.call_count, calls)
        # Once on upload, once on lazy rebuild; switching roles never re-embeds.
        self.assertEqual(self.embeddings.document_calls, 2)

    def test_restart_rebuilds_only_on_first_company_search(self):
        original = b"Original annual leave policy"
        saved = self.policies.upload("c_101", "leave.txt", original)
        restarted = service.PolicyService(Path(self.directory.name), lambda: self.embeddings)
        self.assertEqual(restarted.list("c_101")[0]["id"], saved["id"])
        self.assertEqual(self.embeddings.document_calls, 1)
        restarted.search("c_101", "leave")
        restarted.search("c_101", "leave again")
        self.assertEqual(self.embeddings.document_calls, 2)
        with closing(sqlite3.connect(restarted.database)) as db:
            self.assertEqual(db.execute("SELECT data FROM policies").fetchone()[0], original)
            columns = [row[1] for row in db.execute("PRAGMA table_info(policies)")]
            self.assertNotIn("vector", columns)

    def test_deletion_removes_disk_bytes_and_cached_vectors_without_reembedding(self):
        first = self.policies.upload("c_101", "first.txt", b"First policy")
        second = self.policies.upload("c_101", "second.txt", b"Second policy")
        calls = self.embeddings.document_calls
        self.policies.delete(first["id"])
        self.assertEqual([p["id"] for p in self.policies.list("c_101")], [second["id"]])
        self.assertTrue(all(d.metadata["policy_id"] == second["id"] for d in self.policies.search("c_101", "policy")))
        self.assertEqual(self.embeddings.document_calls, calls)
        self.policies.delete(second["id"])
        self.assertEqual(self.policies.search("c_101", "policy"), [])
        self.assertEqual(self.embeddings.document_calls, calls)
        with self.assertRaises(service.PolicyNotFound):
            self.policies.delete(first["id"])

    def test_delete_before_lazy_rebuild_does_not_load_model(self):
        saved = self.policies.upload("c_101", "leave.txt", b"Policy")
        restarted = service.PolicyService(Path(self.directory.name), lambda: self.embeddings)
        restarted.delete(saved["id"])
        self.assertEqual(restarted.search("c_101", "leave"), [])
        self.assertEqual(self.embeddings.document_calls, 1)

    def test_embedding_failure_does_not_commit_file_or_change_cached_search(self):
        saved = self.policies.upload("c_101", "original.txt", b"Original policy")
        self.embeddings.fail = True
        with self.assertRaises(RuntimeError):
            self.policies.upload("c_101", "failed.txt", b"Failed policy")
        self.assertEqual([p["id"] for p in self.policies.list("c_101")], [saved["id"]])
        self.assertEqual(len(self.policies.search("c_101", "original")), 1)

    def test_parallel_duplicate_upload_is_idempotent(self):
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(
                lambda _: self.policies.upload("c_101", "same.txt", b"Same policy"), range(2)
            ))
        self.assertEqual(results[0]["id"], results[1]["id"])
        self.assertEqual(self.embeddings.document_calls, 1)
        self.assertEqual(len(self.policies.list("c_101")), 1)

    def test_frontend_http_contract_all_formats_and_validation(self):
        http = TestClient(app)
        with patch.object(routes, "get_service", return_value=self.policies):
            self.assertEqual(http.get("/policies", params={"company_id": "c_101"}).json(), [])
            for name, data in (("a.pdf", pdf_bytes()), ("b.docx", docx_bytes()), ("c.txt", b"Text policy")):
                response = http.post("/policies", data={"company_id": "c_101"}, files={"file": (name, data)})
                self.assertEqual(response.status_code, 201)
                self.assertEqual(set(response.json()), {"id", "filename", "size", "uploaded_at"})
            policies = http.get("/policies", params={"company_id": "c_101"}).json()
            self.assertEqual(len(policies), 3)
            self.assertEqual(http.delete("/policies/" + policies[0]["id"]).status_code, 204)
            self.assertEqual(http.delete("/policies/" + policies[0]["id"]).status_code, 404)
            self.assertEqual(http.get("/policies", params={"company_id": "unknown"}).status_code, 404)
            self.assertEqual(http.post("/policies", data={"company_id": "unknown"}, files={"file": ("a.txt", b"hi")}).status_code, 404)
            self.assertEqual(http.post("/policies", data={"company_id": "c_101"}, files={"file": ("a.txt", b"\xff")}).status_code, 422)
            self.assertEqual(http.get("/policies").status_code, 422)


if __name__ == "__main__":
    unittest.main()
